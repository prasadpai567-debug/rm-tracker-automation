"""Native Monthly Analysis dashboard; all business results link to Pivot Summary."""

from copy import deepcopy
from datetime import datetime

from openpyxl.chart import BarChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import AxDataSource, NumData, NumVal, StrData, StrRef, StrVal
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.fill import GradientFillProperties, GradientStop, LinearShadeProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor
from openpyxl.drawing.text import CharacterProperties, Font as DrawingFont
from openpyxl.drawing.text import Paragraph, ParagraphProperties
from openpyxl.drawing.xdr import XDRPoint2D, XDRPositiveSize2D
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.views import Selection

PALETTE = ("4472C4", "C0504D", "9BBB59")


def _text(size=10, bold=False, color="D9D9D9"):
    props = CharacterProperties(
        sz=int(size * 100), b=bold, solidFill=color,
        latin=DrawingFont(typeface="Calibri"),
    )
    return RichText(p=[Paragraph(
        pPr=ParagraphProperties(defRPr=props), endParaRPr=deepcopy(props)
    )])


def _gradient(*colors):
    return GradientFillProperties(
        gsLst=[GradientStop(pos=round(i * 100000 / (len(colors) - 1)), srgbClr=color)
               for i, color in enumerate(colors)],
        lin=LinearShadeProperties(ang=0, scaled=True),
    )


def _shade(color, amount):
    channels = [int(color[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(value + (255 - value) * amount) if amount >= 0 else round(value * (1 + amount)):02X}"
                   for value in channels)


def _paint_bars(series, color):
    # A light centre and darker edges reproduce the reference's bevel-like bars
    # while keeping the chart two dimensional and editable in Excel.
    series.graphicalProperties = GraphicalProperties(
        gradFill=_gradient(_shade(color, -.3), _shade(color, .3), color),
        ln=LineProperties(solidFill=color, w=9525),
    )
    series.invertIfNegative = False


def _style_chart(chart, title, left):
    chart.title = title
    chart.style = None  # No theme-dependent defaults for text or labels.
    chart.roundedCorners = True
    chart.visible_cells_only = False
    chart.graphical_properties = GraphicalProperties(
        gradFill=_gradient("292929", "555555", "292929"),
        ln=LineProperties(noFill=True),
    )
    chart.plot_area.spPr = GraphicalProperties(
        noFill=True, ln=LineProperties(noFill=True)
    )
    chart.layout = Layout(manualLayout=ManualLayout(
        layoutTarget="outer", xMode="edge", yMode="edge",
        x=0.10, y=0.20, w=0.84, h=0.61,
    ))
    chart.title.overlay = False
    chart.title.layout = None  # Excel centres the title within each chart panel.
    chart.title.txPr = _text(16, True, "F2F2F2")
    chart.title.txPr.p[0].pPr.algn = "ctr"
    for para in chart.title.tx.rich.p:
        para.pPr = _text(16, True, "F2F2F2").p[0].pPr
        para.pPr.algn = "ctr"
        para.endParaRPr = deepcopy(para.pPr.defRPr)
        for run in para.r:
            run.rPr = deepcopy(para.pPr.defRPr)
    chart.legend.position = "b"
    chart.legend.overlay = False
    chart.legend.layout = Layout(manualLayout=ManualLayout(
        xMode="edge", yMode="edge", x=0.15, y=0.91, w=0.70, h=0.06,
    ))
    chart.legend.txPr = _text(10)
    for axis in (chart.x_axis, chart.y_axis):
        axis.delete = False
        axis.txPr = _text(9)
        axis.spPr = GraphicalProperties(
            ln=LineProperties(solidFill="9A9A9A", w=9525)
        )
        axis.majorTickMark = "none"
        axis.minorTickMark = "none"
        if axis.title:
            axis.title.txPr = _text(9, True)
            for para in axis.title.tx.rich.p:
                for run in para.r:
                    run.rPr = CharacterProperties(
                        sz=900, b=True, solidFill="D9D9D9",
                        latin=DrawingFont(typeface="Calibri"),
                    )
    chart.y_axis.majorGridlines = ChartLines(
        spPr=GraphicalProperties(ln=LineProperties(solidFill="666666", w=6350))
    )
    chart.dLbls = DataLabelList(
        showLegendKey=False, showVal=True, showCatName=False,
        showSerName=False, showPercent=False, showBubbleSize=False,
        showLeaderLines=False, numFmt="0.##",
        dLblPos="outEnd", txPr=_text(9),
    )
    chart.anchor = AbsoluteAnchor(
        pos=XDRPoint2D(x=pixels_to_EMU(left), y=pixels_to_EMU(255)),
        ext=XDRPositiveSize2D(cx=pixels_to_EMU(675), cy=pixels_to_EMU(338)),
    )


def _initial_values(pivot, combined):
    """Seed chart caches only; Excel keeps the live references on recalculation."""
    import pandas as pd

    names = [str(pivot.cell(row, 1).value or "") for row in range(5, 8)]
    months = [str(pivot.cell(2, col).value) for col in range(2, 5)]
    year = int(pivot["B1"].value)
    dates = pd.to_datetime(combined["PO date"], errors="coerce")
    regions = combined["Region"].fillna("").astype(str).str.strip()
    present_po = combined["PO#"].notna() & combined["PO#"].astype(str).str.strip().ne("")
    counts = []
    for month in months:
        month_number = datetime.strptime(month, "%B").month
        period_mask = (dates.dt.year == year) & (dates.dt.month == month_number)
        counts.append([
            int((period_mask & regions.eq(name) & present_po).sum()) for name in names
        ])
    changes = [
        (counts[2][i] - counts[0][i]) / counts[0][i] * 100 if counts[0][i] else 0
        for i in range(len(names))
    ]
    return names, months, counts, changes


def _cache(series, values, categories, category_formula):
    series.val.numRef.numCache = NumData(
        formatCode="General", ptCount=len(values),
        pt=[NumVal(idx=i, v=value) for i, value in enumerate(values)],
    )
    series.cat = AxDataSource(strRef=StrRef(
        f=category_formula,
        strCache=StrData(
            ptCount=len(categories),
            pt=[StrVal(idx=i, v=value) for i, value in enumerate(categories)],
        ),
    ))


def write_dashboard(wb, combined):
    ws = wb["Monthly Analysis"]
    pivot = wb["Pivot Summary"]
    names, months, counts, changes = _initial_values(pivot, combined)

    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    ws.delete_rows(1, max(ws.max_row, 1))
    ws.column_dimensions.clear()
    ws.row_dimensions.clear()
    ws._charts = []
    ws.conditional_formatting._cf_rules.clear()
    ws.data_validations.dataValidation = []
    ws.auto_filter.ref = None
    ws.freeze_panes = None
    # Removing frozen panes must also discard their pane-specific selections.
    # Otherwise Excel repairs the worksheet view and may group sheet tabs.
    ws.sheet_view.selection = [Selection(activeCell="A1", sqref="A1")]
    for sheet in wb.worksheets:
        sheet.sheet_view.tabSelected = sheet is ws
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.sheet_view.topLeftCell = "A1"
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    for col in "ABCDEFGHIJKLMNOPQRSTU":
        ws.column_dimensions[col].width = 10.7
    for col, width in {
        "D": 25, "E": 17.3, "F": 17.3, "G": 14, "H": 17.3, "I": 16, "J": 16
    }.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 30
    for row in range(4, 8):
        ws.row_dimensions[row].height = 22
    for row in range(8, 27):
        ws.row_dimensions[row].height = 18

    ws.merge_cells("D1:J1")
    ws["D1"] = "Regional PO Monthly Trend Analysis"
    ws["D1"].font = Font(name="Calibri", size=20, bold=True, color="0B2745")
    ws["D1"].alignment = Alignment(horizontal="center", vertical="center")
    for col, text in zip("DEFGHIJ", [
        "Category", "Month 1", "Month 2", "Month 3", "Trend", "Change", "Status"
    ]):
        ws[f"{col}3"] = text
    for col, pivot_col in zip("EFG", "BCD"):
        ws[f"{col}3"] = f"='Pivot Summary'!{pivot_col}4"
    for row, pivot_row in zip(range(4, 7), range(5, 8)):
        for col, pivot_col in zip("DEFGHIJ", "ABCDEFG"):
            ws[f"{col}{row}"] = f"='Pivot Summary'!{pivot_col}{pivot_row}"
    ws["D7"] = "Total POs"
    for col in "EFG":
        ws[f"{col}7"] = f"=SUM({col}4:{col}6)"
    ws["H7"] = '=IF(AND(G7>F7,G7>E7),"Growing",IF(AND(G7<F7,G7<E7),"Declining",IF(AND(G7=F7,G7=E7),"Stable","Volatile")))'
    ws["I7"] = "=IFERROR((G7-E7)/E7*100,0)"
    ws["J7"] = '=IF(I7>10,"Strong",IF(I7>0,"Good",IF(I7=0,"Flat","Down")))'

    line = Side(style="thin", color="000000")
    for row in ws.iter_rows(min_row=3, max_row=7, min_col=4, max_col=10):
        for cell in row:
            cell.border = Border(left=line, right=line, top=line, bottom=line)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(name="Calibri", size=14, color="0B2745", bold=cell.row == 7)
            cell.fill = PatternFill("solid", fgColor="FFFFFF")
            if cell.row == 3:
                cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="244E7D")
    for row in range(4, 8):
        for col in "EFG":
            ws[f"{col}{row}"].number_format = "#,##0"
        ws[f"I{row}"].number_format = "0.00"
    for status, color in {
        "Strong": "FF00B050", "Good": "FFFFC000", "Flat": "FFD9D9D9", "Down": "FFFF0000"
    }.items():
        ws.conditional_formatting.add("J4:J10000", FormulaRule(
            formula=[f'TRIM($J4)="{status}"'],
            fill=PatternFill("solid", fgColor=color, bgColor=color),
            font=Font(color="FF000000"), stopIfTrue=True,
        ))

    volume = BarChart()
    volume.type = "col"
    volume.grouping = "clustered"
    volume.overlap = 0
    volume.gapWidth = 80
    volume.y_axis.title = "PO COUNT"
    volume.x_axis.title = "REGION"
    volume.y_axis.scaling.min = 0
    volume.y_axis.numFmt = "0"
    volume.add_data(Reference(ws, min_col=5, max_col=7, min_row=3, max_row=6),
                    titles_from_data=True)
    for series, color, values, month in zip(volume.series, PALETTE, counts, months):
        _paint_bars(series, color)
        _cache(series, values, names, "'Monthly Analysis'!$D$4:$D$6")
        series.tx.strRef.strCache = StrData(ptCount=1, pt=[StrVal(idx=0, v=month)])
    _style_chart(volume, "Regional PO Volume by Month", 58)
    ws.add_chart(volume)

    change = BarChart()
    change.type = "bar"
    change.grouping = "clustered"
    change.varyColors = True
    change.gapWidth = 100
    change.x_axis.axPos = "l"
    change.y_axis.axPos = "b"
    change.y_axis.title = "CHANGE (%)"
    change.y_axis.numFmt = "0.##"
    # Direct visible ranges: no hidden helper column can suppress these bars.
    change.add_data(Reference(ws, min_col=9, max_col=9, min_row=3, max_row=6),
                    titles_from_data=True)
    _cache(change.series[0], changes, names, "'Monthly Analysis'!$D$4:$D$6")
    change.series[0].invertIfNegative = False
    for i, color in enumerate(PALETTE):
        point = DataPoint(idx=i, invertIfNegative=False)
        point.graphicalProperties = GraphicalProperties(
            gradFill=_gradient(_shade(color, -.3), _shade(color, .3), color),
            ln=LineProperties(solidFill=color, w=9525),
        )
        change.series[0].data_points.append(point)
    _style_chart(change, "Three-Month Change by Region", 756)
    ws.add_chart(change)
    ws.print_options.horizontalCentered = True
    ws.print_area = "A1:R26"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    wb.active = wb.index(ws)
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
