"""Audit and approved-defect regressions; no real Outlook connection.

All workbook writes are temporary. Run from the project root with unittest
discovery. Approved defects must pass, not remain expected failures.
"""
from __future__ import annotations
import hashlib
import os
import shutil
import sys
import tempfile
import types
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import MagicMock, patch

import pandas as pd
from openpyxl import Workbook, load_workbook
import rm_tracker_app as app

PROJECT = Path(__file__).resolve().parents[1]
AUDIT = PROJECT / "audit_20261028"
SOURCE = AUDIT / "inputs" / "RM_Followup_Dummy_Data.xlsx"
CONTACTS = AUDIT / "inputs" / "Vendor Contact List.xlsx"
BASELINE = AUDIT / "baseline" / "RM_Tracker_Pivot_Analysis_Final.xlsx"
COLORS = {"Strong": "FF00B050", "Good": "FFFFC000", "Flat": "FFD9D9D9", "Down": "FFFF0000"}
AUDIT_FIXTURES_AVAILABLE = all(path.exists() for path in (SOURCE, CONTACTS, BASELINE))


def history_frame():
    return pd.DataFrame([{"Follow-up Type": "No Confirmation", "PO#": "P1",
        "Item Number": "10", "Vendor Number": "100", "Supplier email id": "test@example.invalid",
        "Reminder": "Reminder 1", "Created Date": datetime.combine(date.today(), datetime.min.time()),
        "Email Status": "Sent"}])


def follow_frame():
    return pd.DataFrame([{"Follow-up type": "No Confirmation", "PO#": "P1",
        "Item Number": "10", "Vendor Number": "100", "Supplier email id": "test@example.invalid",
        "Reminder": "Reminder 1", "Manager email": "manager@example.invalid"}])


class FakeOutlook:
    def __init__(self, fail_at=None):
        self.messages = []
        self.fail_at = fail_at
    def GetNamespace(self, name):
        return self
    def GetDefaultFolder(self, value):
        return types.SimpleNamespace(Name="TEST ONLY")
    def CreateItem(self, value):
        owner = self
        class Message:
            def Save(self):
                if owner.fail_at == len(owner.messages):
                    raise RuntimeError("simulated interruption")
                owner.messages.append(self)
            def Send(self):
                raise AssertionError("Sending is forbidden in the audit")
        return Message()


def mock_outlook(fake):
    package = types.ModuleType("win32com")
    client = types.ModuleType("win32com.client")
    client.Dispatch = lambda name: fake
    package.client = client
    return patch.dict(sys.modules, {"win32com": package, "win32com.client": client})


@unittest.skipUnless(
    AUDIT_FIXTURES_AVAILABLE,
    "Private audit workbook fixtures are not included in the source repository.",
)
class AuditTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.combined = app.read_tracker(SOURCE)
        cls.baseline = load_workbook(BASELINE)

    @classmethod
    def tearDownClass(cls):
        cls.baseline.close()

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="case_", dir=AUDIT)
        self.folder = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)

    def write_tracker(self, frame, name="tracker.xlsx"):
        path = self.folder / name
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            for region in app.TRACKER_SHEETS:
                frame.to_excel(writer, sheet_name=region, index=False)
        return path

    def generate(self, output, source=SOURCE, drafts=False):
        return app.generate(source, CONTACTS, output, drafts)

    def test_import_row_and_region_counts(self):
        raw = {name: pd.read_excel(SOURCE, sheet_name=name) for name in app.TRACKER_SHEETS}
        self.assertEqual(len(self.combined), sum(len(f) for f in raw.values()))
        self.assertEqual(self.combined["Source Sheet"].value_counts().to_dict(),
                         {name: len(f) for name, f in raw.items()})

    def test_tracker_metadata_handle_is_closed(self):
        book=MagicMock()
        book.sheet_names=app.TRACKER_SHEETS
        book.__enter__.return_value=book
        with patch.object(app.pd,"ExcelFile",return_value=book), patch.object(
                app.pd,"read_excel",return_value=self.combined.iloc[:1].copy()):
            app.read_tracker(SOURCE)
        self.assertTrue(book.close.called or book.__exit__.called,
                        "Tracker metadata reader must close its workbook handle")

    def test_manager_metadata_handle_is_closed(self):
        book=MagicMock()
        book.sheet_names=["contacts"]
        book.__enter__.return_value=book
        with patch.object(app.pd,"ExcelFile",return_value=book), patch.object(
                app.pd,"read_excel",return_value=pd.DataFrame({"Vendor Number":["1"]})):
            app.read_manager_contacts(CONTACTS)
        self.assertTrue(book.close.called or book.__exit__.called,
                        "Manager metadata reader must close its workbook handle")

    def test_source_identifiers_material_serials_and_headers_preserved(self):
        for region in app.TRACKER_SHEETS:
            raw = app.clean_columns(pd.read_excel(SOURCE, sheet_name=region))
            got = self.combined[self.combined["Source Sheet"].eq(region)].reset_index(drop=True)
            self.assertEqual(got.columns.tolist(), raw.columns.tolist() + ["Source Sheet"])
            for col in ("Material Number", "Slno", "Vendor Name", "Qty", "Unit price"):
                pd.testing.assert_series_equal(got[col], raw[col], check_dtype=False, check_names=False)
            for col in ("Vendor Number", "PO#"):
                self.assertEqual(got[col].tolist(), raw[col].map(app.normalise_vendor).tolist())

    def test_no_new_duplicates_or_deduplication(self):
        raw = pd.concat([pd.read_excel(SOURCE, sheet_name=s) for s in app.TRACKER_SHEETS])
        self.assertEqual(int(self.combined.duplicated().sum()), int(raw.duplicated().sum()))
        frame = self.combined.iloc[[0, 0]].copy()
        path = self.write_tracker(frame)
        self.assertEqual(len(app.read_tracker(path)), 6)

    def test_combined_sheet_matches_import_exactly(self):
        actual = pd.read_excel(BASELINE, sheet_name="Combined_Data")
        expected = self.combined.copy()
        for col in ("Vendor Number", "PO#"):
            actual[col] = actual[col].map(app.normalise_vendor)
        pd.testing.assert_frame_equal(actual, expected, check_dtype=False)

    def test_excel_dates_are_typed(self):
        ws = self.baseline["Combined_Data"]
        for header in ("PO date", "Delivery date", "Delivery Date2"):
            col = next(c.column for c in ws[1] if c.value == header)
            for row in range(2, ws.max_row + 1):
                c = ws.cell(row, col)
                if c.value is not None:
                    self.assertIsInstance(c.value, datetime)
                    self.assertTrue(c.is_date)

    def test_followup_eligibility_independent_oracle(self):
        expected_nc, expected_d = [], []
        today = date.today()
        for _, row in self.combined.iterrows():
            status = app.text(row["Status"]).casefold()
            po_date = row["PO date"]
            delivery_date = row["Delivery Date2"]
            if status == "no confirmation" and pd.notna(po_date) and po_date.date() < today:
                expected_nc.append(row["PO#"])
            if (status == "complete" and "delivered" not in app.text(row["Comments"]).casefold()
                    and pd.notna(delivery_date) and delivery_date.date() <= today):
                expected_d.append(row["PO#"])
        for kind, expected in (("no_confirmation", expected_nc), ("delivery", expected_d)):
            actual = app.build_followups(self.combined, {}, kind)
            self.assertEqual(actual["PO#"].tolist(), expected)

    def test_no_confirmation_date_status_boundaries(self):
        today = pd.Timestamp(date.today())
        frame = pd.concat([self.combined.iloc[[0]]] * 6, ignore_index=True)
        frame["Status"] = ["No Confirmation", "no CONFIRMATION ", "No Confirmation",
                           "No Confirmation", "Deleted", "Complete"]
        frame["PO date"] = [today-timedelta(days=1), today-timedelta(days=2),
                            today, today+timedelta(days=1), today-timedelta(days=2), pd.NaT]
        frame["PO#"] = [str(n) for n in range(6)]
        self.assertEqual(app.build_followups(frame, {}, "no_confirmation")["PO#"].tolist(), ["0", "1"])

    def test_delivery_comments_status_and_date_boundaries(self):
        today = pd.Timestamp(date.today())
        frame = pd.concat([self.combined.iloc[[0]]] * 8, ignore_index=True)
        frame["Status"] = ["Complete", "complete ", "Complete", "Complete", "Deleted",
                           "Partially delivered", "Complete", "Complete"]
        frame["Comments"] = ["", None, "Delivered", "not delivered", "", "", "", ""]
        frame["Delivery Date2"] = [today, today-timedelta(days=1), today, today, today,
                                   today, today+timedelta(days=1), pd.NaT]
        frame["PO#"] = [str(n) for n in range(8)]
        self.assertEqual(app.build_followups(frame, {}, "delivery")["PO#"].tolist(), ["0", "1"])

    def test_contact_mapping_combines_unique_addresses(self):
        path = self.folder / "contacts.xlsx"
        pd.DataFrame({"Vendor Number": [100,100,100,None],
                      "Email Address": ["one@example.invalid","one@example.invalid","two@example.invalid","bad"]}).to_excel(path,index=False)
        self.assertEqual(app.read_contacts(path,"Sheet1"),
                         {"100": "one@example.invalid; two@example.invalid"})

    def test_blank_vendor_does_not_invent_contact(self):
        frame = self.combined.iloc[[0]].copy()
        frame["Vendor Number"] = None
        frame["Status"] = "No Confirmation"
        frame["PO date"] = pd.Timestamp(date.today()-timedelta(days=1))
        row = app.build_followups(frame, {"100":"someone@example.invalid"}, "no_confirmation").iloc[0]
        self.assertEqual(row["Vendor Number"], "")
        self.assertEqual(row["Supplier email id"], "")

    def test_parse_dates_invalid_and_excel_serial(self):
        parsed = app.parse_excel_dates(pd.Series([45000, "invalid", None]))
        self.assertEqual(parsed.iloc[0], pd.Timestamp("2023-03-15"))
        self.assertTrue(pd.isna(parsed.iloc[1]))
        self.assertTrue(pd.isna(parsed.iloc[2]))

    def test_invalid_import_dates_are_coerced_to_missing(self):
        frame = self.combined.iloc[[0]].copy()
        frame["PO date"] = "invalid"
        actual = app.read_tracker(self.write_tracker(frame))
        self.assertTrue(actual["PO date"].isna().all())

    def test_all_required_sheets_formulas_and_recalculation(self):
        wb = self.baseline
        self.assertEqual(wb.sheetnames, app.OUTPUT_SHEETS)
        self.assertEqual(wb.calculation.calcMode, "auto")
        self.assertTrue(wb.calculation.fullCalcOnLoad)
        self.assertTrue(wb.calculation.forceFullCalc)
        self.assertFalse(any(ws._pivots for ws in wb))
        for ws in wb:
            for row in ws:
                for c in row:
                    if c.data_type == "f":
                        self.assertNotIn("#REF!", c.value)

    def test_pivot_native_dropdowns_and_no_autofilter(self):
        ws = self.baseline["Pivot Summary"]
        self.assertIsNone(ws.auto_filter.ref)
        validations = ws.data_validations.dataValidation
        for addr in ("B1", "B2", "C2", "D2"):
            self.assertEqual(sum(addr in dv for dv in validations), 1)
        months = next(dv for dv in validations if "B2" in dv)
        self.assertEqual(months.formula1.strip('"').split(","),
            ["January","February","March","April","May","June","July","August","September","October","November","December"])

    def test_pivot_formulas_and_source_columns_normal_schema(self):
        ws = self.baseline["Pivot Summary"]
        headers = [c.value for c in self.baseline["Combined_Data"][1]]
        self.assertEqual((headers[21], headers[6], headers[9]), ("Region","PO#","PO date"))
        for row in range(5, 8):
            self.assertEqual(ws[f"F{row}"].value, f"=IFERROR((D{row}-B{row})/B{row}*100,0)")
            self.assertEqual(ws[f"G{row}"].value,
                f'=IF(F{row}>10,"Strong",IF(F{row}>0,"Good",IF(F{row}=0,"Flat","Down")))')
            self.assertIn(f'D{row}>C{row},D{row}>B{row}', ws[f"E{row}"].value)

    def test_regional_monthly_summary_counts_independent(self):
        actual = app.make_monthly_analysis(self.combined)
        expected = {}
        for _, row in self.combined.iterrows():
            key = (row["Region"], row["PO date"].strftime("%Y-%m"))
            expected[key] = expected.get(key, 0) + (0 if pd.isna(row["PO#"]) else 1)
        self.assertEqual({(r.Region,r.Month): getattr(r,"_2") for r in actual.itertuples(index=False)}, expected)

    def test_status_rules_saved_argb_solid_black_and_data_only(self):
        for sheet, col, start in (("Pivot Summary","G",5),("Monthly Analysis","J",4)):
            ws = self.baseline[sheet]
            entries = [(key,rules) for key,rules in ws.conditional_formatting._cf_rules.items()
                       if str(key.sqref) == f"{col}{start}:{col}10000"]
            self.assertEqual(len(entries),1)
            rules=entries[0][1]
            self.assertEqual(len(rules),4)
            for rule in rules:
                status = rule.formula[0].split('"')[1]
                self.assertEqual(rule.formula, [f'TRIM(${col}{start})="{status}"'])
                self.assertEqual(rule.dxf.fill.fgColor.rgb, COLORS[status])
                self.assertEqual(rule.dxf.fill.fill_type, "solid")
                self.assertEqual(rule.dxf.font.color.rgb, "FF000000")

    def test_dashboard_links_totals_and_chart_references(self):
        ws=self.baseline["Monthly Analysis"]
        for row, source_row in zip(range(4,7),range(5,8)):
            for col, source_col in zip("DEFGHIJ","ABCDEFG"):
                self.assertEqual(ws[f"{col}{row}"].value,f"='Pivot Summary'!{source_col}{source_row}")
        self.assertEqual(ws["E7"].value,"=SUM(E4:E6)")
        self.assertEqual(len(ws._charts),2)
        for chart in ws._charts:
            self.assertFalse(chart.visible_cells_only)
            for s in chart.series:
                self.assertIsNotNone(s.val.numRef)
                self.assertIsNotNone(s.cat.strRef)
                self.assertIn("'Monthly Analysis'!",s.val.numRef.f)

    def test_repeat_generation_all_sheets_and_formulas_identical(self):
        path=self.folder/"RM_Tracker_Pivot_Analysis_Final.xlsx"
        self.generate(path)
        first=load_workbook(path)
        before={s.title:[list(row) for row in s.values] for s in first}
        first.close()
        self.generate(path)
        second=load_workbook(path)
        self.assertEqual(before,{s.title:[list(row) for row in s.values] for s in second})
        self.assertEqual(len(second["Monthly Analysis"]._charts),2)
        second.close()

    def test_generation_without_drafts_preserves_history(self):
        path=self.folder/"RM_Tracker_Pivot_Analysis_Final.xlsx"
        shutil.copy2(AUDIT/"backup"/"RM_Tracker_Pivot_Analysis_Final.xlsx",path)
        before=pd.read_excel(path,sheet_name="Follow_Up_History")
        self.generate(path)
        after=pd.read_excel(path,sheet_name="Follow_Up_History")
        pd.testing.assert_frame_equal(before,after)

    def test_successful_followup_count_sent_only_and_item_specific(self):
        frame=pd.concat([history_frame()]*4,ignore_index=True)
        frame["Email Status"]=["Sent","Email sent","Draft saved","send confirmed"]
        frame.loc[3,"Item Number"]="20"
        self.assertEqual(app.successful_followup_count(frame,"no confirmation","P1","10"),3)

    def test_mock_draft_body_cc_save_no_send(self):
        fake=FakeOutlook()
        history=pd.concat([history_frame()]*2,ignore_index=True)
        history["Created Date"]=[
            pd.Timestamp(date.today()-timedelta(days=2)),
            pd.Timestamp(date.today()-timedelta(days=1)),
        ]
        with mock_outlook(fake):
            result=app.create_email_drafts([follow_frame()],history,self.folder)
        self.assertEqual(result,(1,0,0))
        self.assertEqual(fake.messages[0].To,"test@example.invalid")
        self.assertEqual(fake.messages[0].CC,"manager@example.invalid")
        self.assertIn("P1",fake.messages[0].Subject)

    def test_mock_draft_skips_same_day_and_missing_contact(self):
        fake=FakeOutlook()
        frame=pd.concat([follow_frame()]*2,ignore_index=True)
        frame.loc[1,"Supplier email id"]=""
        with mock_outlook(fake):
            result=app.create_email_drafts([frame],history_frame(),self.folder)
        self.assertEqual(result,(1,1,0))
        self.assertEqual(len(fake.messages),1)

    def test_history_driven_reminders_advance_and_same_day_reuses(self):
        history=pd.concat([history_frame()]*2,ignore_index=True)
        history["Email Status"]="Draft saved"
        history["Reminder"]=["Reminder 1","Reminder 2"]
        history["Created Date"]=[
            pd.Timestamp(date.today()-timedelta(days=2)),
            pd.Timestamp(date.today()-timedelta(days=1)),
        ]
        due=app.apply_history_reminders(follow_frame(),history)
        self.assertEqual(due.iloc[0]["Reminder"],"Reminder 3")

        today=history_frame()
        today["Email Status"]="Draft saved"
        today["Reminder"]="Reminder 3"
        history=pd.concat([history,today],ignore_index=True)
        rerun=app.apply_history_reminders(follow_frame(),history)
        self.assertEqual(rerun.iloc[0]["Reminder"],"Reminder 3")

    def test_same_day_reminder_two_rerun_does_not_escalate(self):
        history=pd.concat([history_frame()]*2,ignore_index=True)
        history["Email Status"]="Draft saved"
        history["Reminder"]=["Reminder 1","Reminder 2"]
        history["Created Date"]=[
            pd.Timestamp(date.today()-timedelta(days=1)),
            pd.Timestamp(date.today()),
        ]
        fake=FakeOutlook()
        with mock_outlook(fake):
            result=app.create_email_drafts([follow_frame()],history,self.folder)
        self.assertEqual(result,(1,0,0))
        self.assertFalse(hasattr(fake.messages[0],"CC"))

    def test_mock_draft_interruption_has_descriptive_error(self):
        with mock_outlook(FakeOutlook(fail_at=0)):
            with self.assertRaisesRegex(RuntimeError,"Outlook could not save the draft"):
                app.create_email_drafts([follow_frame()],pd.DataFrame(),self.folder)

    def test_missing_input_missing_sheet_and_missing_required_columns(self):
        with self.assertRaises(FileNotFoundError):
            self.generate(self.folder/"out.xlsx",self.folder/"missing.xlsx")
        bad=self.folder/"bad.xlsx"
        pd.DataFrame({"Wrong":[1]}).to_excel(bad,index=False)
        with self.assertRaisesRegex(ValueError,"missing required sheets"):
            app.read_tracker(bad)
        with self.assertRaisesRegex(ValueError,"must contain Vendor Number"):
            app.read_contacts(bad,"Sheet1")
        frame=self.combined.iloc[[0]].drop(columns=["Status"])
        with self.assertRaises(KeyError):
            self.generate(self.folder/"out.xlsx",self.write_tracker(frame))

    def test_empty_file_and_empty_valid_tracker(self):
        invalid=self.folder/"empty.xlsx"
        invalid.touch()
        with self.assertRaises(ValueError):
            app.read_tracker(invalid)
        path=self.write_tracker(self.combined.iloc[:0])
        result=self.generate(self.folder/"out.xlsx",path)
        self.assertEqual(result["rows"],0)
        self.assertEqual(result["no_confirmation"],0)
        self.assertEqual(result["delivery"],0)

    def test_output_folder_created_and_locked_file_fails(self):
        result=self.generate(self.folder/"new"/"folder"/"out.xlsx")
        self.assertTrue(Path(result["output"]).exists())
        if os.name!="nt":
            return
        import ctypes
        from ctypes import wintypes
        path=self.folder/"locked.xlsx"
        shutil.copy2(BASELINE,path)
        kernel=ctypes.WinDLL("kernel32",use_last_error=True)
        kernel.CreateFileW.argtypes=[wintypes.LPCWSTR,wintypes.DWORD,wintypes.DWORD,
                                    ctypes.c_void_p,wintypes.DWORD,wintypes.DWORD,wintypes.HANDLE]
        kernel.CreateFileW.restype=wintypes.HANDLE
        kernel.CloseHandle.argtypes=[wintypes.HANDLE]
        handle=kernel.CreateFileW(str(path),0x80000000,0,None,3,0x80,None)
        self.assertNotEqual(handle,ctypes.c_void_p(-1).value)
        try:
            with self.assertRaises(PermissionError):
                self.generate(path)
        finally:
            kernel.CloseHandle(handle)

    def test_known_missing_region_fallback_crashes(self):
        frame=self.combined.iloc[[0]].drop(columns=["Region"])
        self.assertEqual(len(app.read_tracker(self.write_tracker(frame))),3)

    def test_known_reordered_source_breaks_formula_column_mapping(self):
        frame=self.combined.iloc[[0]].copy()
        frame=frame[["Region"]+[c for c in frame if c!="Region"]]
        path=self.folder/"out.xlsx"
        self.generate(path,self.write_tracker(frame))
        wb=load_workbook(path)
        try:
            region_column=next(c.column_letter for c in wb["Combined_Data"][1] if c.value=="Region")
            self.assertIn(f"Combined_Data!${region_column}$2:",wb["Pivot Summary"]["B5"].value)
        finally:
            wb.close()

    def test_known_import_numeric_dates_become_1970(self):
        frame=self.combined.iloc[[0]].copy()
        frame["PO date"]=45000
        actual=app.read_tracker(self.write_tracker(frame))
        self.assertEqual(actual["PO date"].iloc[0],pd.Timestamp("2023-03-15"))

    def test_known_regeneration_loses_user_month_selection(self):
        path=self.folder/"out.xlsx"
        self.generate(path)
        wb=load_workbook(path)
        wb["Pivot Summary"]["B1"]=2026
        wb["Pivot Summary"]["B2"]="March"
        wb["Pivot Summary"]["C2"]="April"
        wb["Pivot Summary"]["D2"]="May"
        wb.save(path)
        wb.close()
        self.generate(path)
        wb=load_workbook(path)
        try:
            self.assertEqual([wb["Pivot Summary"][f"{c}2"].value for c in "BCD"],["March","April","May"])
        finally:
            wb.close()

    def test_known_history_read_failure_silently_returns_empty(self):
        path=self.folder/"unreadable.xlsx"
        path.write_bytes(b"not a workbook")
        with self.assertRaises(Exception):
            app.existing_history(path)

    def test_known_history_same_day_overwrites_old_timestamp_status(self):
        old=history_frame()
        with patch.object(app,"existing_history",return_value=old.copy()):
            actual=app.history_rows([follow_frame()],self.folder/"unused.xlsx","No draft created")
        pd.testing.assert_frame_equal(actual.reset_index(drop=True),old,check_dtype=False)

    def test_same_day_success_upgrades_failed_history_without_duplicate(self):
        old=history_frame()
        old["Email Status"]="No draft created"
        drafted={app.followup_key("No Confirmation","P1","10")}
        with patch.object(app,"existing_history",return_value=old.copy()):
            actual=app.history_rows(
                [follow_frame()],self.folder/"unused.xlsx","Draft saved",drafted
            )
        self.assertEqual(len(actual),1)
        self.assertEqual(actual.iloc[0]["Email Status"],"Draft saved")
        self.assertEqual(actual.iloc[0]["Reminder"],"Reminder 1")

    def test_known_history_marks_skipped_contact_as_draft_saved(self):
        path=self.folder/"out.xlsx"
        contacts={kind:app.read_contacts(CONTACTS,sheet) for kind,sheet in app.CONTACT_SHEETS.items()}
        nc=app.build_followups(self.combined,contacts["no_confirmation"],"no_confirmation")
        delivery=app.build_followups(self.combined,contacts["delivery"],"delivery")
        missing_rows=pd.concat([nc,delivery]).loc[lambda f:f["Supplier email id"].eq("")]
        missing={
            (row["Follow-up type"],row["PO#"],row["Item Number"])
            for _,row in missing_rows.iterrows()
        }
        self.assertTrue(missing)
        fake=FakeOutlook()
        with mock_outlook(fake):
            self.generate(path,drafts=True)
        hist=pd.read_excel(path,sheet_name="Follow_Up_History",dtype={"PO#":str})
        skipped=hist.loc[hist.apply(
            lambda row:(row["Follow-up Type"],row["PO#"],row["Item Number"]) in missing,
            axis=1,
        )]
        self.assertFalse(skipped["Email Status"].eq("Draft saved").any())

    def test_known_save_failure_leaves_partial_existing_output(self):
        from openpyxl.workbook.workbook import Workbook
        path=self.folder/"out.xlsx"
        shutil.copy2(BASELINE,path)
        before=hashlib.sha256(path.read_bytes()).hexdigest()
        real_save=Workbook.save
        calls=0
        def fail_second(wb,*args,**kwargs):
            nonlocal calls
            calls+=1
            if calls==2:
                raise OSError("simulated final-save interruption")
            return real_save(wb,*args,**kwargs)
        with patch.object(Workbook,"save",fail_second):
            with self.assertRaisesRegex(OSError,"simulated"):
                self.generate(path)
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),before)

    def test_known_duplicate_drafts_with_same_key_in_one_run(self):
        frame=pd.concat([follow_frame()]*2,ignore_index=True)
        fake=FakeOutlook()
        with mock_outlook(fake):
            app.create_email_drafts([frame],pd.DataFrame(),self.folder)
        self.assertEqual(len(fake.messages),1)

    def test_mixed_draft_outcomes_are_logged_per_key(self):
        frame=pd.concat([follow_frame()]*3,ignore_index=True)
        frame.loc[0,"Supplier email id"]=""
        frame.loc[2,"PO#"]="P2"
        frame.loc[2,"Supplier email id"]=""
        drafted=set()
        fake=FakeOutlook()
        with mock_outlook(fake):
            counts=app.create_email_drafts([frame],pd.DataFrame(),self.folder,drafted)
        self.assertEqual(counts,(1,2,0))
        self.assertEqual(len(fake.messages),1)
        history=app.history_rows([frame],self.folder/"new.xlsx","Draft saved",drafted)
        self.assertEqual(len(history),2)
        self.assertEqual(dict(zip(history["PO#"],history["Email Status"])),
                         {"P1":"Draft saved","P2":"No draft created"})

    def test_existing_duplicate_history_rows_are_not_erased(self):
        old=pd.concat([history_frame()]*2,ignore_index=True)
        old.loc[1,"Created Date"]=old.loc[0,"Created Date"]+timedelta(hours=1)
        with patch.object(app,"existing_history",return_value=old.copy()):
            actual=app.history_rows([follow_frame()],self.folder/"unused.xlsx","No draft created")
        pd.testing.assert_frame_equal(actual.reset_index(drop=True),old,check_dtype=False)

    def test_corrupt_existing_output_is_never_replaced(self):
        path=self.folder/"corrupt.xlsx"
        original=b"not a workbook"
        path.write_bytes(original)
        with self.assertRaises(Exception):
            self.generate(path)
        self.assertEqual(path.read_bytes(),original)

    def test_interrupted_first_save_leaves_no_new_output(self):
        from openpyxl.workbook.workbook import Workbook
        path=self.folder/"new.xlsx"
        real_save=Workbook.save
        calls=0
        def fail_second(wb,*args,**kwargs):
            nonlocal calls
            calls+=1
            if calls==2:
                raise OSError("simulated final-save interruption")
            return real_save(wb,*args,**kwargs)
        with patch.object(Workbook,"save",fail_second):
            with self.assertRaisesRegex(OSError,"simulated"):
                self.generate(path)
        self.assertFalse(path.exists())
        self.assertEqual(list(self.folder.glob(".new.*.xlsx")),[])

    def test_selected_future_year_and_month_order_survive_regeneration(self):
        path=self.folder/"out.xlsx"
        self.generate(path)
        wb=load_workbook(path)
        wb["Pivot Summary"]["B1"]=2027
        for column,month in zip("BCD",["December","February","January"]):
            wb["Pivot Summary"][f"{column}2"]=month
        wb.save(path)
        wb.close()
        self.generate(path)
        wb=load_workbook(path)
        try:
            pivot=wb["Pivot Summary"]
            self.assertEqual(pivot["B1"].value,2027)
            self.assertEqual([pivot[f"{c}2"].value for c in "BCD"],
                             ["December","February","January"])
            self.assertEqual(len(wb["Monthly Analysis"]._charts),2)
        finally:
            wb.close()


if __name__=="__main__":
    unittest.main(verbosity=2)
