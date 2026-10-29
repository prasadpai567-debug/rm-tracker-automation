"""Regression checks for chart data visibility, OOXML and sheet isolation."""
import io
import unittest
import zipfile
import xml.etree.ElementTree as ET

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.views import Selection
from monthly_dashboard import write_dashboard


class DashboardTests(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        self.wb.active.title = "Combined_Data"
        self.pivot = self.wb.create_sheet("Pivot Summary")
        self.ws = self.wb.create_sheet("Monthly Analysis")
        self.pivot["B1"] = 2026
        for col, month in zip("BCD", ("August", "September", "October")):
            self.pivot[f"{col}2"] = month
            self.pivot[f"{col}4"] = f"={col}2"
        records = []
        expected = [(32, 43, 17), (41, 52, 20), (30, 23, 17)]
        for row, name, counts in zip(range(5, 8), ("Blair", "Franklinton", "Salem"), expected):
            self.pivot[f"A{row}"] = name
            for month, count in zip((8, 9, 10), counts):
                records.extend({"Region": name, "PO date": pd.Timestamp(2026, month, 1),
                                "PO#": f"{name}-{month}-{n}"} for n in range(count))
        self.combined = pd.DataFrame(records)
        self.wb["Combined_Data"]["A1"] = "preserve this"
        self.wb["Combined_Data"]["A2"] = "=1+2"
        self.ws.freeze_panes = "B4"
        self.ws.sheet_view.selection = [Selection(pane="bottomRight", sqref="A1:G1")]
        self.pivot.sheet_view.tabSelected = True

    def build(self):
        write_dashboard(self.wb, self.combined)
        data = io.BytesIO()
        self.wb.save(data)
        return data.getvalue()

    def test_direct_live_data_and_caches(self):
        self.build()
        self.assertEqual(self.ws["E4"].value, "='Pivot Summary'!B5")
        self.assertEqual(self.ws["J4"].value, "='Pivot Summary'!G5")
        self.assertEqual(len(self.ws._charts), 2)
        volume, change = self.ws._charts
        self.assertEqual([pt.v for pt in volume.series[0].val.numRef.numCache.pt], [32, 41, 30])
        self.assertEqual(change.series[0].val.numRef.f, "'Monthly Analysis'!$I$4:$I$6")
        self.assertEqual(len(change.series[0].data_points), 3)
        self.assertFalse(change.visible_cells_only)
        self.assertFalse(any(v.hidden for v in self.ws.column_dimensions.values()))

    def test_valid_sheet_view_after_removing_freeze_panes(self):
        content = self.build()
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            root = ET.fromstring(archive.read("xl/worksheets/sheet3.xml"))
            ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            selections = root.findall("s:sheetViews/s:sheetView/s:selection", ns)
            self.assertEqual(len(selections), 1)
            self.assertNotIn("pane", selections[0].attrib)
        self.assertTrue(self.ws.sheet_view.tabSelected)
        self.assertFalse(self.pivot.sheet_view.tabSelected)

    def test_chart_xml_has_background_and_unambiguous_labels(self):
        content = self.build()
        ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
              "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            for name in ("xl/charts/chart1.xml", "xl/charts/chart2.xml"):
                root = ET.fromstring(archive.read(name))
                self.assertIsNotNone(root.find("c:spPr/a:gradFill", ns))
                for stop in root.findall(".//a:gs", ns):
                    self.assertTrue(stop.get("pos").isdigit())
                    self.assertLessEqual(int(stop.get("pos")), 100000)
                self.assertEqual(root.find(".//c:dLbls/c:showSerName", ns).get("val"), "0")
                self.assertEqual(root.find(".//c:dLbls/c:showCatName", ns).get("val"), "0")
                self.assertEqual(root.find("c:chart/c:title/c:overlay", ns).get("val"), "0")

    def test_other_content_and_all_status_rules_are_preserved(self):
        content = self.build()
        reopened = load_workbook(io.BytesIO(content))
        self.assertEqual(reopened["Combined_Data"]["A2"].value, "=1+2")
        self.assertEqual(reopened["Pivot Summary"]["B2"].value, "August")
        rules = list(reopened["Monthly Analysis"].conditional_formatting._cf_rules.values())[0]
        colors = {rule.formula[0]: rule.dxf.fill.fgColor.rgb for rule in rules}
        self.assertEqual(colors, {
            'TRIM($J4)="Strong"': "FF00B050", 'TRIM($J4)="Good"': "FFFFC000",
            'TRIM($J4)="Flat"': "FFD9D9D9", 'TRIM($J4)="Down"': "FFFF0000",
        })


if __name__ == "__main__":
    unittest.main()
