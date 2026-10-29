from __future__ import annotations

import argparse
import logging
from html import escape
import os
import re
import sys
import tempfile
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

import pandas as pd


LOGGER = logging.getLogger("rm_tracker_app")


TRACKER_SHEETS = ["Franklinton", "Blair", "Salem"]
CONTACT_SHEETS = {
    "no_confirmation": "No Confirmation_Follow_UP",
    "delivery": "Delivery FolloW_Up",
}
MANAGER_EMAIL_COLUMNS = (
    "Manager Email",
    "Manager Email Address",
    "Manager Email ID",
    "Responsible Manager Email",
)
OUTPUT_SHEETS = [
    "Combined_Data",
    "Pivot Summary",
    "Monthly Analysis",
    "NO Confirmation Follow_UP",
    "Delivery followUP",
    "Follow_Up_History",
]

SUCCESSFUL_EMAIL_STATUSES = frozenset(
    {"sent", "email sent", "send confirmed", "draft saved"}
)


def clean_columns(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    return df


def text(value) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip()


def normalise_vendor(value) -> str:
    value = text(value)
    if value.endswith(".0"):
        value = value[:-2]
    return value


def format_date(value) -> str:
    if pd.isna(value) or value == "":
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    parsed = pd.to_datetime(value, errors="coerce")
    return "" if pd.isna(parsed) else parsed.strftime("%Y-%m-%d")


def read_tracker(path: Path) -> pd.DataFrame:
    with pd.ExcelFile(path) as workbook:
        sheets = workbook.sheet_names
    missing = [sheet for sheet in TRACKER_SHEETS if sheet not in sheets]
    if missing:
        raise ValueError(f"Tracker is missing required sheets: {', '.join(missing)}")

    frames = []
    for sheet in TRACKER_SHEETS:
        frame = clean_columns(pd.read_excel(path, sheet_name=sheet))
        if "Region" in frame:
            frame["Region"] = frame["Region"].fillna(sheet).replace("", sheet)
        else:
            frame["Region"] = sheet
        frame["Source Sheet"] = sheet
        frames.append(frame)
    combined = pd.concat(frames, ignore_index=True)
    for column in ["Vendor Number", "PO#"]:
        if column in combined:
            combined[column] = combined[column].map(normalise_vendor)
    for column in ["PO date", "Delivery date", "Delivery Date2"]:
        if column in combined:
            combined[column] = parse_excel_dates(combined[column])
    return combined


def read_contacts(path: Path, sheet_name: str) -> dict[str, str]:
    frame = clean_columns(pd.read_excel(path, sheet_name=sheet_name))
    required = {"Vendor Number", "Email Address"}
    if not required.issubset(frame.columns):
        raise ValueError(f"{sheet_name} must contain Vendor Number and Email Address columns")
    mapping: dict[str, list[str]] = defaultdict(list)
    for _, row in frame.iterrows():
        vendor = normalise_vendor(row["Vendor Number"])
        email = text(row["Email Address"])
        if vendor and email and email not in mapping[vendor]:
            mapping[vendor].append(email)
    return {vendor: "; ".join(emails) for vendor, emails in mapping.items()}


def read_manager_contacts(path: Path) -> dict[str, str]:
    """Read optional manager mappings without inventing a fallback address."""
    result: dict[str, str] = {}
    with pd.ExcelFile(path) as workbook:
        sheet_names = workbook.sheet_names
    for sheet_name in sheet_names:
        frame = clean_columns(pd.read_excel(path, sheet_name=sheet_name))
        manager_column = next((column for column in MANAGER_EMAIL_COLUMNS if column in frame.columns), None)
        if not manager_column or "Vendor Number" not in frame.columns:
            continue
        for _, row in frame.iterrows():
            vendor = normalise_vendor(row.get("Vendor Number"))
            email = text(row.get(manager_column))
            if vendor and email and "@" in email and vendor not in result:
                result[vendor] = email
    return result


def item_number(row: pd.Series) -> str:
    for column in ("Item Number", "PO Item", "Line Item", "Item"):
        if column in row.index:
            value = text(row.get(column))
            if value:
                return value
    return ""


def reminder_number(value: str) -> int:
    match = re.search(r"(?:second|2nd|reminder\s*2|follow[- ]?up\s*2)", value, re.I)
    if match:
        return 2
    match = re.search(r"(?:third|3rd|reminder\s*3|follow[- ]?up\s*3)", value, re.I)
    if match:
        return 3
    return 1


def parse_excel_dates(values: pd.Series) -> pd.Series:
    """Parse Excel date values and date strings into comparable timestamps."""
    parsed = pd.to_datetime(values, errors="coerce")
    numeric = pd.to_numeric(values, errors="coerce")
    excel_serials = numeric.notna() & numeric.between(1, 2958465)
    if excel_serials.any():
        parsed.loc[excel_serials] = pd.to_datetime(
            numeric.loc[excel_serials], unit="D", origin="1899-12-30", errors="coerce"
        )
    return parsed


def build_followups(
    combined: pd.DataFrame,
    contacts: dict[str, str],
    kind: str,
    managers: dict[str, str] | None = None,
) -> pd.DataFrame:
    status = combined["Status"].map(text).str.casefold()
    if kind == "no_confirmation":
        # Include only overdue no-confirmation POs; today's POs are excluded.
        po_dates = parse_excel_dates(combined["PO date"])
        rows = combined[
            (status == "no confirmation")
            & po_dates.notna()
            & (po_dates.dt.normalize() < pd.Timestamp(date.today()))
        ].copy()
        comment_column = "OC Followup comments"
        followup_type = "No Confirmation"
    else:
        # Delivery follow-up requires a completed PO, no Delivered comment, and
        # a delivery date that is today or earlier.
        comments = combined.get("Comments", pd.Series("", index=combined.index)).map(text)
        delivery_dates = parse_excel_dates(combined["Delivery Date2"])
        rows = combined[
            (status == "complete")
            & ~comments.str.contains("Delivered", case=False, na=False)
            & delivery_dates.notna()
            & (delivery_dates.dt.normalize() <= pd.Timestamp(date.today()))
        ].copy()
        comment_column = "Delivery Followups Comments"
        followup_type = "Delivery"

    managers = managers or {}
    records = []
    for _, row in rows.iterrows():
        vendor = normalise_vendor(row.get("Vendor Number"))
        comment = text(row.get(comment_column, ""))
        records.append(
            {
                "Vendor Number": vendor,
                "PO#": normalise_vendor(row.get("PO#")),
                "Supplier email id": contacts.get(vendor, ""),
                "Reminder": f"Reminder {reminder_number(comment)}",
                "Vendor Name": text(row.get("Vendor Name")),
                "Region": text(row.get("Region")),
                "PO date": row.get("PO date"),
                "Expected delivery date": row.get("Delivery date"),
                "Order Confirmation": text(row.get("Order Confirmation")),
                "Follow-up type": followup_type,
                "Notes": comment,
                "Item Number": item_number(row),
                "Manager email": managers.get(vendor, ""),
            }
        )
    return pd.DataFrame(records)


def month_label(series: pd.Series) -> pd.Series:
    parsed = pd.to_datetime(series, errors="coerce")
    return parsed.dt.strftime("%Y-%m")


def make_monthly_analysis(combined: pd.DataFrame) -> pd.DataFrame:
    work = combined.copy()
    work["Month"] = month_label(work["PO date"])
    work["PO Value"] = pd.to_numeric(work["Qty"], errors="coerce").fillna(0) * pd.to_numeric(
        work["Unit price"], errors="coerce"
    ).fillna(0)
    work["No Confirmation"] = work["Order Confirmation"].isna().astype(int)
    work["Delivery Follow-up"] = (
        work["Order Confirmation"].notna() & work["Delivery Date2"].isna()
    ).astype(int)
    result = (
        work.groupby(["Region", "Month"], dropna=False)
        .agg(
            **{
                "PO Count": ("PO#", "count"),
                "Total Quantity": ("Qty", "sum"),
                "PO Value": ("PO Value", "sum"),
                "No Confirmation": ("No Confirmation", "sum"),
                "Delivery Follow-up": ("Delivery Follow-up", "sum"),
            }
        )
        .reset_index()
        .sort_values(["Month", "Region"])
    )
    return result


def make_pivot_summary(combined: pd.DataFrame) -> pd.DataFrame:
    """Build the monthly region trend layout used by the meeting workbook."""
    work = combined.copy()
    work["_PODate"] = parse_excel_dates(work["PO date"])
    work = work[work["_PODate"].notna()].copy()
    if work.empty:
        return pd.DataFrame(columns=["Category", "Category", "Trend", "Change", "Status"])

    work["_Month"] = work["_PODate"].dt.to_period("M")
    monthly = (
        work.groupby(["Region", "_Month"])["PO#"]
        .count()
        .unstack(fill_value=0)
        .sort_index(axis=1)
    )
    periods = list(monthly.columns[-3:])
    month_labels = [period.strftime("%B") for period in periods]
    records = []

    for region, values in monthly.iterrows():
        counts = [int(values.get(period, 0)) for period in periods]
        first = counts[0] if counts else 0
        last = counts[-1] if counts else 0
        change = ((last - first) / first * 100) if first else (100.0 if last else 0.0)
        differences = [counts[i] - counts[i - 1] for i in range(1, len(counts))]
        if not differences or all(diff == 0 for diff in differences):
            trend = "Stable"
        elif last > first and any(diff < 0 for diff in differences):
            trend = "Volatile"
        elif last < first and any(diff > 0 for diff in differences):
            trend = "Volatile"
        elif last > first:
            trend = "Growing"
        else:
            trend = "Declining"
        status = "Strong" if change > 0 else ("Down" if change < 0 else "Flat")
        records.append(
            ["", region, *counts, trend, change, status]
        )

    columns = ["Category", "Category", *month_labels, "Trend", "Change", "Status"]
    result = pd.DataFrame(records, columns=columns)
    order = result.iloc[:, 1].astype(str).str.casefold().argsort(kind="stable")
    return result.iloc[order].reset_index(drop=True)


def write_monthly_analysis_dashboard(
    wb,
    combined: pd.DataFrame,
    preserved_selection: tuple[object, list[object]] | None,
) -> None:
    """Build the linked view after the Pivot Summary formulas have been written."""
    from monthly_dashboard import write_dashboard

    write_dashboard(wb, combined)

def existing_history(output_path: Path) -> pd.DataFrame:
    columns = ["Follow-up Type", "PO#", "Item Number", "Vendor Number", "Supplier email id", "Reminder", "Created Date", "Email Status"]
    if not output_path.exists():
        return pd.DataFrame(columns=columns)
    return clean_columns(pd.read_excel(output_path, sheet_name="Follow_Up_History"))


def successful_history_matches(
    history: pd.DataFrame, followup_type: str, po: str, item: str
) -> pd.DataFrame:
    """Return completed follow-ups for one PO line, including saved drafts."""
    if history.empty:
        return history.iloc[0:0].copy()
    work = history.copy()
    for column in (
        "Follow-up Type", "PO#", "Item Number", "Reminder", "Created Date", "Email Status"
    ):
        if column not in work:
            work[column] = ""
    completed = work[
        work["Email Status"].map(text).str.casefold().isin(SUCCESSFUL_EMAIL_STATUSES)
    ]
    return completed[
        (completed["Follow-up Type"].map(text).str.casefold() == followup_type.casefold())
        & (completed["PO#"].map(normalise_vendor) == normalise_vendor(po))
        & (completed["Item Number"].map(text) == text(item))
    ]


def successful_followup_count(
    history: pd.DataFrame, followup_type: str, po: str, item: str
) -> int:
    return len(successful_history_matches(history, followup_type, po, item))


def history_reminder_number(
    history: pd.DataFrame, followup_type: str, po: str, item: str
) -> int | None:
    """Return the reminder due now, reusing today's reminder on correction reruns."""
    matches = successful_history_matches(history, followup_type, po, item)
    if matches.empty:
        return None

    created = pd.to_datetime(matches["Created Date"], errors="coerce")
    today_mask = created.dt.date == date.today()
    if today_mask.any():
        today_rows = matches.loc[today_mask]
        for value in reversed(today_rows["Reminder"].tolist()):
            if text(value):
                return reminder_number(value)

    prior_dates = {
        value.date()
        for value in created[created.notna()]
        if value.date() < date.today()
    }
    undated_count = int(created.isna().sum())
    return min(max(len(prior_dates) + undated_count + 1, 1), 3)


def apply_history_reminders(frame: pd.DataFrame, history: pd.DataFrame) -> pd.DataFrame:
    """Apply history-driven reminders without changing the source follow-up rows."""
    result = frame.copy()
    if result.empty:
        return result
    for index, row in result.iterrows():
        number = history_reminder_number(
            history,
            text(row.get("Follow-up type")),
            text(row.get("PO#")),
            text(row.get("Item Number", "")),
        )
        if number is not None:
            result.at[index, "Reminder"] = f"Reminder {number}"
    return result


def followup_key(kind: object, po: object, item: object) -> tuple[str, str, str]:
    return text(kind).casefold(), normalise_vendor(po), text(item)


def history_rows(
    followups: Iterable[pd.DataFrame],
    output_path: Path,
    status: str,
    drafted_keys: set[tuple[str, str, str]] | None = None,
) -> pd.DataFrame:
    old = existing_history(output_path).copy()
    today_indices = {}
    for index, row in old.iterrows():
        created = pd.to_datetime(row.get("Created Date"), errors="coerce")
        if pd.notna(created) and created.date() == date.today():
            key = followup_key(
                row.get("Follow-up Type"), row.get("PO#"), row.get("Item Number")
            )
            today_indices[key] = index
    new_by_key = {}
    for frame in followups:
        for _, row in frame.iterrows():
            key = followup_key(
                row["Follow-up type"], row["PO#"], row.get("Item Number", "")
            )
            key_today = (*key, date.today())
            row_status = (
                "Draft saved" if key in drafted_keys else "No draft created"
            ) if drafted_keys is not None else status
            new_row = {
                "Follow-up Type": row["Follow-up type"],
                "PO#": row["PO#"],
                "Item Number": row.get("Item Number", ""),
                "Vendor Number": row["Vendor Number"],
                "Supplier email id": row["Supplier email id"],
                "Reminder": row["Reminder"],
                "Created Date": datetime.now(),
                "Email Status": row_status,
            }
            if key in today_indices:
                old_index = today_indices[key]
                old_status = text(old.at[old_index, "Email Status"]).casefold()
                if (
                    row_status == "Draft saved"
                    and old_status not in SUCCESSFUL_EMAIL_STATUSES
                ):
                    for column, value in new_row.items():
                        old.at[old_index, column] = value
                continue
            if key_today not in new_by_key or (
                row_status == "Draft saved"
                and new_by_key[key_today]["Email Status"] != "Draft saved"
            ):
                new_by_key[key_today] = new_row
    if not new_by_key:
        return old
    return pd.concat([old, pd.DataFrame(new_by_key.values())], ignore_index=True)


def _write_workbook_file(
    output_path: Path,
    previous_output_path: Path,
    combined: pd.DataFrame,
    pivot: pd.DataFrame,
    monthly: pd.DataFrame,
    no_confirmation: pd.DataFrame,
    delivery: pd.DataFrame,
    history: pd.DataFrame,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    output_path.parent.mkdir(parents=True, exist_ok=True)
    preserved_selection = (None, [])
    if previous_output_path.exists():
        previous = load_workbook(previous_output_path, read_only=True, data_only=False)
        try:
            if "Pivot Summary" in previous:
                previous_ws = previous["Pivot Summary"]
                preserved_selection = (
                    previous_ws["B1"].value,
                    [previous_ws[f"{column}2"].value for column in "BCD"],
                )
        finally:
            previous.close()
    visible_followup_columns = [
        "Vendor Number", "PO#", "Supplier email id", "Reminder", "Vendor Name",
        "Region", "PO date", "Expected delivery date", "Order Confirmation",
        "Follow-up type", "Notes",
    ]
    no_confirmation = no_confirmation.reindex(columns=[c for c in visible_followup_columns if c in no_confirmation.columns])
    delivery = delivery.reindex(columns=[c for c in visible_followup_columns if c in delivery.columns])
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        combined.to_excel(writer, sheet_name="Combined_Data", index=False)
        pivot.to_excel(writer, sheet_name="Pivot Summary", index=False, startrow=3)
        monthly.to_excel(writer, sheet_name="Monthly Analysis", index=False)
        no_confirmation.to_excel(writer, sheet_name="NO Confirmation Follow_UP", index=False)
        delivery.to_excel(writer, sheet_name="Delivery followUP", index=False)
        history.to_excel(writer, sheet_name="Follow_Up_History", index=False)

    wb = load_workbook(output_path)
    header_fill = PatternFill("solid", fgColor="0F6B78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows():
            for cell in row:
                if cell.column in (1, 2, 3, 4, 5, 6) and cell.value is not None:
                    cell.alignment = Alignment(vertical="top", wrap_text=False)
        for column_cells in ws.columns:
            letter = get_column_letter(column_cells[0].column)
            max_len = max(len(str(cell.value or "")) for cell in column_cells[:200])
            ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 32)
    pivot_ws = wb["Pivot Summary"]
    # Rebuild only this generated sheet so every future workbook has the same
    # formula-driven 3-month summary, regardless of the source file's size.
    pivot_ws.delete_rows(1, max(pivot_ws.max_row, 1))
    pivot_ws.sheet_view.showGridLines = False
    date_values = parse_excel_dates(combined["PO date"])
    periods = sorted(date_values.dropna().dt.to_period("M").unique())
    if periods:
        selected_periods = list(periods[-3:])
    else:
        current = pd.Timestamp(date.today()).to_period("M")
        selected_periods = [current - 2, current - 1, current]
    while len(selected_periods) < 3:
        selected_periods.insert(0, selected_periods[0] - 1)
    selected_months = [period.strftime("%B") for period in selected_periods]
    previous_year, previous_months = preserved_selection
    valid_months = set(
        "January,February,March,April,May,June,July,August,September,October,November,December".split(",")
    )
    if len(previous_months) == 3 and all(month in valid_months for month in previous_months):
        selected_months = list(previous_months)
    years = sorted(set(date_values.dropna().dt.year.astype(int).tolist()))
    if not years:
        years = [date.today().year]
    selected_year = max(years)
    if (
        isinstance(previous_year, (int, float))
        and not isinstance(previous_year, bool)
        and pd.notna(previous_year)
        and int(previous_year) == previous_year
        and 1900 <= int(previous_year) <= 9999
    ):
        selected_year = int(previous_year)
    year_values = sorted(
        set(years) | {max(years) + offset for offset in range(1, 4)} | {selected_year}
    )
    regions = [str(value).strip() for value in combined["Region"].dropna().unique() if str(value).strip()]
    regions = sorted(dict.fromkeys(regions), key=str.casefold)
    if not regions:
        regions = ["Blair", "Franklinton", "Salem"]

    pivot_ws["A1"] = "Year"
    pivot_ws["B1"] = selected_year
    pivot_ws["A2"] = "Select Months"
    pivot_ws["B2"] = selected_months[0]
    pivot_ws["C2"] = selected_months[1]
    pivot_ws["D2"] = selected_months[2]
    pivot_ws["A3"] = "Count of PO#"
    pivot_ws["B3"] = "Count of PO#"
    pivot_ws["C3"] = "Column Labels"
    for address in ("E2", "F2", "G2", "C1", "D1", "E1", "F1", "G1", "D3", "E3", "F3", "G3"):
        pivot_ws[address].value = None
    for address, value in {
        "A4": "Category", "B4": "=B2", "C4": "=C2", "D4": "=D2",
        "E4": "Trend", "F4": "Change", "G4": "Status",
    }.items():
        pivot_ws[address] = value

    source_last_row = len(combined) + 1
    region_column = get_column_letter(combined.columns.get_loc("Region") + 1)
    po_column = get_column_letter(combined.columns.get_loc("PO#") + 1)
    date_column = get_column_letter(combined.columns.get_loc("PO date") + 1)
    for row_number, region_name in enumerate(regions, start=5):
        pivot_ws.cell(row=row_number, column=1).value = region_name
        for column, header_ref in zip(("B", "C", "D"), ("B$4", "C$4", "D$4")):
            pivot_ws[f"{column}{row_number}"] = (
                f'=COUNTIFS(Combined_Data!${region_column}$2:${region_column}${source_last_row},$A{row_number},'
                f'Combined_Data!${po_column}$2:${po_column}${source_last_row},"<>",'
                f'Combined_Data!${date_column}$2:${date_column}${source_last_row},">="&DATE($B$1,'
                f'MONTH(DATEVALUE({header_ref}&" 1")),1),'
                f'Combined_Data!${date_column}$2:${date_column}${source_last_row},"<"&EDATE(DATE($B$1,'
                f'MONTH(DATEVALUE({header_ref}&" 1")),1),1))'
            )
        pivot_ws[f"E{row_number}"] = (
            f'=IF(AND(D{row_number}>C{row_number},D{row_number}>B{row_number}),"Growing",'
            f'IF(AND(D{row_number}<C{row_number},D{row_number}<B{row_number}),"Declining",'
            f'IF(AND(D{row_number}=C{row_number},D{row_number}=B{row_number}),"Stable","Volatile")))'
        )
        pivot_ws[f"F{row_number}"] = f'=IFERROR((D{row_number}-B{row_number})/B{row_number}*100,0)'
        pivot_ws[f"G{row_number}"] = (
            f'=IF(F{row_number}>10,"Strong",IF(F{row_number}>0,"Good",'
            f'IF(F{row_number}=0,"Flat","Down")))'
        )

    year_validation = DataValidation(
        type="list", formula1='"' + ",".join(str(year) for year in year_values) + '"', allow_blank=False
    )
    month_validation = DataValidation(
        type="list",
        formula1='"January,February,March,April,May,June,July,August,September,October,November,December"',
        allow_blank=False,
    )
    pivot_ws.add_data_validation(year_validation)
    year_validation.add(pivot_ws["B1"])
    pivot_ws.add_data_validation(month_validation)
    for address in ("B2", "C2", "D2"):
        month_validation.add(pivot_ws[address])

    dark = "17324D"
    navy = "0B2745"
    border_color = "B7C9D6"
    thin = Side(style="thin", color=border_color)
    for row in pivot_ws.iter_rows(min_row=1, max_row=4 + len(regions), min_col=1, max_col=7):
        for cell in row:
            cell.font = Font(name="Calibri", size=11, color=dark)
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for address in ("A1", "A2", "A3", "B3", "C3"):
        pivot_ws[address].font = Font(name="Calibri", size=11, bold=True, color=dark)
        pivot_ws[address].fill = PatternFill("solid", fgColor="D9EAF7")
    for address in ("B1", "B2", "C2", "D2"):
        pivot_ws[address].font = Font(name="Calibri", size=11, bold=True, color=dark)
        pivot_ws[address].fill = PatternFill("solid", fgColor="FFF2CC")
        pivot_ws[address].alignment = Alignment(horizontal="center", vertical="center")
    for address in ("A4", "B4", "C4", "D4", "E4", "F4", "G4"):
        pivot_ws[address].font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        pivot_ws[address].fill = PatternFill("solid", fgColor=navy)
        pivot_ws[address].alignment = Alignment(horizontal="center", vertical="center")
    last_summary_row = 4 + len(regions)
    for row_number in range(5, last_summary_row + 1):
        pivot_ws[f"A{row_number}"].font = Font(name="Calibri", size=11, bold=True, color=dark)
        for column in "BCD":
            pivot_ws[f"{column}{row_number}"].number_format = "0"
            pivot_ws[f"{column}{row_number}"].alignment = Alignment(horizontal="right")
        pivot_ws[f"F{row_number}"].number_format = "0.00"
        pivot_ws[f"F{row_number}"].alignment = Alignment(horizontal="right")
        for column in "EG":
            pivot_ws[f"{column}{row_number}"].alignment = Alignment(horizontal="center")
    pivot_ws.conditional_formatting.add(
        f"F5:F{last_summary_row}", CellIsRule(operator="greaterThan", formula=["10"], fill=PatternFill("solid", fgColor="C6EFCE"))
    )
    pivot_ws.conditional_formatting.add(
        f"F5:F{last_summary_row}", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor="FFC7CE"))
    )
    # Status-column rules are formula-based because Status cells contain
    # formulas. Locate the header instead of assuming a fixed column, remove
    # only existing rules for that column, and then add native Excel rules.
    status_header = next(
        (cell for row in pivot_ws.iter_rows() for cell in row
         if str(cell.value).strip().casefold() == "status"),
        None,
    )
    if status_header is None:
        raise ValueError("Pivot Summary Status header was not generated")
    status_column = get_column_letter(status_header.column)
    status_data_start = status_header.row + 1
    status_data_end = 10000
    for conditional_range in list(pivot_ws.conditional_formatting._cf_rules):
        if any(
            cell_range.min_col == status_header.column
            and cell_range.max_col == status_header.column
            for cell_range in conditional_range.sqref.ranges
        ):
            del pivot_ws.conditional_formatting._cf_rules[conditional_range]
    status_rules = {
        "Strong": "FF00B050",
        "Good": "FFFFC000",
        "Flat": "FFD9D9D9",
        "Down": "FFFF0000",
    }
    for status_text, fill_color in status_rules.items():
        pivot_ws.conditional_formatting.add(
            f"{status_column}{status_data_start}:{status_column}{status_data_end}",
            FormulaRule(
                formula=[f'TRIM($' + f'{status_column}{status_data_start})="{status_text}"'],
                fill=PatternFill(fill_type="solid", fgColor=fill_color, bgColor=fill_color),
                font=Font(color="FF000000"),
                stopIfTrue=True,
            ),
        )
    for column, width in {"A": 20, "B": 15, "C": 15, "D": 15, "E": 14, "F": 12, "G": 14}.items():
        pivot_ws.column_dimensions[column].width = width
    pivot_ws.row_dimensions[1].height = 22
    pivot_ws.row_dimensions[2].height = 22
    pivot_ws.row_dimensions[3].height = 22
    pivot_ws.row_dimensions[4].height = 24
    pivot_ws.freeze_panes = "B5"
    pivot_ws.auto_filter.ref = None
    write_monthly_analysis_dashboard(wb, combined, preserved_selection)
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    try:
        wb.save(output_path)
    finally:
        wb.close()


def write_workbook(
    output_path: Path,
    combined: pd.DataFrame,
    pivot: pd.DataFrame,
    monthly: pd.DataFrame,
    no_confirmation: pd.DataFrame,
    delivery: pd.DataFrame,
    history: pd.DataFrame,
) -> None:
    """Replace the final workbook only after its full temporary build succeeds."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, name = tempfile.mkstemp(
        prefix=f".{output_path.stem}.", suffix=".xlsx", dir=output_path.parent
    )
    os.close(descriptor)
    temporary_path = Path(name)
    try:
        _write_workbook_file(
            temporary_path, output_path, combined, pivot, monthly,
            no_confirmation, delivery, history,
        )
        os.replace(temporary_path, output_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def create_email_drafts(
    frames: list[pd.DataFrame],
    history: pd.DataFrame,
    drafts_dir: Path,
    drafted_keys: set[tuple[str, str, str]] | None = None,
) -> tuple[int, int, int]:
    """Create unsent messages in the default Classic Outlook Drafts folder."""
    del drafts_dir  # Retained in the signature for compatibility with callers.
    try:
        import win32com.client
    except ImportError as exc:
        raise RuntimeError(
            "Classic Outlook drafting requires pywin32. Install requirements.txt first."
        ) from exc

    try:
        outlook = win32com.client.Dispatch("Outlook.Application")
        namespace = outlook.GetNamespace("MAPI")
        drafts_folder = namespace.GetDefaultFolder(16)  # olFolderDrafts
    except Exception as exc:
        LOGGER.exception("Unable to connect to the Classic Outlook profile")
        raise RuntimeError(
            "Classic Outlook could not be opened. Start Classic Outlook and confirm that "
            "a default mail profile and Drafts folder are available."
        ) from exc

    history = history.copy()
    for column in ("Follow-up Type", "PO#", "Item Number", "Created Date"):
        if column not in history:
            history[column] = ""
    created = 0
    skipped = 0
    review = 0
    saved_keys = drafted_keys if drafted_keys is not None else set()

    for frame in frames:
        for row_index, (_, row) in enumerate(frame.iterrows(), start=1):
            recipient = text(row["Supplier email id"])
            if not recipient:
                skipped += 1
                continue
            kind = text(row["Follow-up type"])
            po = text(row["PO#"])
            vendor = text(row["Vendor Number"])
            item = text(row.get("Item Number", ""))
            key = followup_key(kind, po, item)
            if key in saved_keys:
                skipped += 1
                continue
            current_reminder = history_reminder_number(history, kind, po, item)

            subject = f"{kind} follow-up for PO {po}"
            manager = text(row.get("Manager email", ""))
            manager_escalation = current_reminder is not None and current_reminder >= 3
            cc = manager if manager_escalation and manager else ""
            if manager_escalation and not manager:
                review += 1
            plain_body = (
                "Dear Supplier,\n\n"
                f"This is a follow-up regarding purchase order {po} "
                f"for vendor number {vendor}.\n\n"
                "Please provide the required update at your earliest convenience.\n\n"
                "Regards"
            )
            html_body = (
                "Dear Supplier,<br><br>"
                f"This is a follow-up regarding purchase order <b>{escape(po)}</b> "
                f"for vendor number <b>{escape(vendor)}</b>.<br><br>"
                "Please provide the required update at your earliest convenience.<br><br>"
                "Regards"
            )
            try:
                message = outlook.CreateItem(0)  # olMailItem
                message.To = recipient
                if cc:
                    message.CC = cc
                message.Subject = subject
                message.Body = plain_body
                message.HTMLBody = html_body
                message.Save()
                saved_keys.add(key)
                created += 1
                LOGGER.info("Saved Outlook draft for PO %s in %s", po, drafts_folder.Name)
            except Exception as exc:
                LOGGER.exception("Unable to save Outlook draft for PO %s", po)
                raise RuntimeError(
                    f"Outlook could not save the draft for PO {po}. "
                    "Confirm that Classic Outlook is running and the Drafts folder is available."
                ) from exc
    return created, skipped, review
def generate(
    tracker_path: Path,
    contacts_path: Path,
    output_path: Path,
    create_drafts: bool = False,
) -> dict:
    combined = read_tracker(tracker_path)
    no_contacts = read_contacts(contacts_path, CONTACT_SHEETS["no_confirmation"])
    delivery_contacts = read_contacts(contacts_path, CONTACT_SHEETS["delivery"])
    managers = read_manager_contacts(contacts_path)
    no_confirmation = build_followups(combined, no_contacts, "no_confirmation", managers)
    delivery = build_followups(combined, delivery_contacts, "delivery", managers)
    pivot = make_pivot_summary(combined)
    monthly = make_monthly_analysis(combined)
    drafts_created = 0
    drafts_skipped = 0
    manager_review = 0
    history_status = "Workbook generated"
    history = existing_history(output_path)
    no_confirmation = apply_history_reminders(no_confirmation, history)
    delivery = apply_history_reminders(delivery, history)
    drafted_keys: set[tuple[str, str, str]] = set()
    if create_drafts:
        drafts_created, drafts_skipped, manager_review = create_email_drafts(
            [no_confirmation, delivery], history,
            output_path.parent / f"{output_path.stem}_Email_Drafts",
            drafted_keys=drafted_keys,
        )
        history_status = "Draft saved" if drafts_created else "No draft created"
    history = history_rows(
        [no_confirmation, delivery], output_path, history_status,
        drafted_keys=drafted_keys,
    ) if create_drafts else history
    write_workbook(output_path, combined, pivot, monthly, no_confirmation, delivery, history)
    return {
        "rows": len(combined),
        "no_confirmation": len(no_confirmation),
        "delivery": len(delivery),
        "drafts_created": drafts_created,
        "drafts_skipped": drafts_skipped,
        "manager_review": manager_review,
        "output": str(output_path),
    }


def _run_gui_legacy() -> None:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    def resource_path(filename: str) -> Path:
        # PyInstaller extracts bundled files under _MEIPASS. In source mode,
        # assets live beside this script.
        bundle_root = getattr(sys, "_MEIPASS", Path(__file__).resolve().parent)
        return Path(bundle_root) / filename

    LOGGER.info("Starting Tkinter GUI")
    try:
        root = tk.Tk()
    except tk.TclError:
        LOGGER.exception("Tkinter could not create the main window")
        raise
    root.title("RM Follow-Up Automation")
    root.geometry("1180x760")
    root.minsize(980, 680)
    root.configure(bg="#F6F8FB")
    try:
        icon_path = resource_path("app_icon.ico")
        if icon_path.exists():
            root.iconbitmap(default=str(icon_path))
    except tk.TclError:
        pass

    colors = {
        "navy": "#0B1F3A", "teal": "#16C7B7", "teal_light": "#E9FBF8",
        "bg": "#F6F8FB", "card": "#FFFFFF", "border": "#DDE5EF",
        "muted": "#6B7A90", "green": "#16A34A", "amber": "#F59E0B",
        "red": "#DC2626",
    }
    style = ttk.Style(root)
    try:
        style.theme_use("clam")
    except tk.TclError:
        pass
    style.configure("Primary.TButton", background=colors["teal"], foreground=colors["navy"],
                    font=("Segoe UI", 10, "bold"), padding=(16, 10), borderwidth=0)
    style.map("Primary.TButton", background=[("active", "#12AFA2"), ("disabled", "#B8DCD8")])
    style.configure("Secondary.TButton", background="#FFFFFF", foreground=colors["navy"],
                    font=("Segoe UI", 10, "bold"), padding=(14, 10), borderwidth=1)
    style.map("Secondary.TButton", background=[("active", colors["teal_light"])])
    style.configure("Link.TButton", background=colors["card"], foreground="#087F78",
                    font=("Segoe UI", 9, "bold"), padding=(6, 4), borderwidth=0)
    style.configure("Modern.Horizontal.TProgressbar", troughcolor="#E7EEF5", background=colors["teal"], borderwidth=0)

    # Keep default data paths beside the project in source mode and beside the
    # packaged application (or its parent project folder) in frozen mode.
    # PyInstaller extracts __file__ under _MEIPASS, which is not the folder
    # containing the user's Vendor Contact List.xlsx.
    runtime_dir = (
        Path(sys.executable).resolve().parent
        if getattr(sys, "frozen", False)
        else Path(__file__).resolve().parent
    )
    project_dir = runtime_dir
    if getattr(sys, "frozen", False):
        packaged_parent = runtime_dir.parent
        if not (runtime_dir / "Vendor Contact List.xlsx").exists() and (
            packaged_parent / "Vendor Contact List.xlsx"
        ).exists():
            project_dir = packaged_parent
    app_dir = project_dir
    default_tracker = app_dir / "RM Tracker.xlsx"
    default_contacts = app_dir / "Vendor Contact List.xlsx"
    default_output = app_dir / "RM_Tracker_Pivot_Analysis_Final.xlsx"
    tracker_var = tk.StringVar(value=str(default_tracker) if default_tracker.exists() else "")
    # The vendor contact list is a stable application input. Keep it connected
    # to the standard file in the application folder so users do not need to
    # browse for it on every run.
    contacts_var = tk.StringVar(value=str(default_contacts))
    output_var = tk.StringVar(value=str(default_output))
    status_title_var = tk.StringVar(value="Ready")
    status_message_var = tk.StringVar(value="Select a tracker workbook, then choose an action.")
    details_var = tk.StringVar(value="")
    dashboard_date_var = tk.StringVar(value=f"As of {date.today():%d %b %Y}")
    no_confirmation_total_var = tk.StringVar(value="—")
    delivery_total_var = tk.StringVar(value="—")
    dashboard_message_var = tk.StringVar(value="Select a valid tracker to calculate daily follow-up totals.")

    def shorten_path(value: str, limit: int = 58) -> str:
        value = value or "No file selected"
        if len(value) <= limit:
            return value
        path = Path(value)
        name = path.name
        return f"{str(path.parent)[:max(12, limit - len(name) - 5)]}...\\{name}"

    def set_status(title: str, message: str, tone: str = "ready"):
        status_title_var.set(title)
        status_message_var.set(message)
        status_dot.configure(bg={"ready": colors["teal"], "success": colors["green"],
                                 "warning": colors["amber"], "error": colors["red"],
                                 "processing": colors["teal"]}.get(tone, colors["teal"]))

    def choose_file(var: tk.StringVar, title: str, save: bool = False):
        if save:
            path = filedialog.asksaveasfilename(
                title=title, defaultextension=".xlsx",
                filetypes=[("Excel workbook", "*.xlsx *.xlsm")],
            )
        else:
            path = filedialog.askopenfilename(
                title=title, filetypes=[("Excel workbook", "*.xlsx *.xlsm")]
            )
        if path:
            var.set(path)

    def refresh_dashboard(*_):
        tracker_path = Path(tracker_var.get()).expanduser() if tracker_var.get() else None
        dashboard_date_var.set(f"As of {date.today():%d %b %Y}")
        if not tracker_path or not tracker_path.exists():
            no_confirmation_total_var.set("—")
            delivery_total_var.set("—")
            dashboard_message_var.set("Select a valid tracker to calculate daily follow-up totals.")
            return
        try:
            combined = read_tracker(tracker_path)
            no_total = len(build_followups(combined, {}, "no_confirmation"))
            delivery_total = len(build_followups(combined, {}, "delivery"))
            no_confirmation_total_var.set(f"{no_total:,}")
            delivery_total_var.set(f"{delivery_total:,}")
            dashboard_message_var.set("Counts include one row per consolidated PO line item.")
        except Exception as exc:
            no_confirmation_total_var.set("—")
            delivery_total_var.set("—")
            dashboard_message_var.set(f"Tracker could not be read: {exc}")

    def open_output_folder():
        target = Path(output_var.get()).expanduser()
        folder = target.parent if target.suffix else target
        if not folder.exists():
            folder = app_dir
        try:
            os.startfile(str(folder))
        except OSError as exc:
            messagebox.showerror("Cannot open folder", str(exc))

    def logo(parent, size=42):
        canvas = tk.Canvas(parent, width=size, height=size, bg=colors["navy"], highlightthickness=0)
        pad = 7
        canvas.create_oval(pad, pad, size - pad, size - pad, outline=colors["teal"], width=2)
        canvas.create_arc(8, 5, size - 5, size - 5, start=35, extent=255, style="arc",
                         outline="#FFFFFF", width=2)
        canvas.create_line(size - 11, 7, size - 5, 10, fill="#FFFFFF", width=2)
        canvas.create_line(size - 11, 7, size - 12, 14, fill="#FFFFFF", width=2)
        canvas.create_text(size / 2, size / 2 + 1, text="RM", fill="#FFFFFF",
                           font=("Segoe UI", max(7, size // 6), "bold"))
        canvas.create_line(size - 17, size - 12, size - 13, size - 8, fill=colors["teal"], width=2)
        canvas.create_line(size - 13, size - 8, size - 7, size - 16, fill=colors["teal"], width=2)
        return canvas

    def process_card(parent, number: str, title: str, description: str, icon: str):
        card = tk.Frame(parent, bg=colors["card"], highlightbackground=colors["border"], highlightthickness=1)
        card.pack(fill="x", pady=(0, 12))
        icon_box = tk.Frame(card, bg=colors["teal_light"], width=38, height=38)
        icon_box.pack(side="left", padx=16, pady=16)
        icon_box.pack_propagate(False)
        tk.Label(icon_box, text=icon, bg=colors["teal_light"], fg="#087F78",
                 font=("Segoe UI Symbol", 15, "bold")).pack(expand=True)
        body = tk.Frame(card, bg=colors["card"])
        body.pack(side="left", fill="x", expand=True, padx=(0, 14), pady=13)
        tk.Label(body, text=f"{number}  {title}", bg=colors["card"], fg=colors["navy"],
                 font=("Segoe UI", 11, "bold"), anchor="w").pack(fill="x")
        tk.Label(body, text=description, bg=colors["card"], fg=colors["muted"],
                 font=("Segoe UI", 9), anchor="w", justify="left", wraplength=360).pack(fill="x", pady=(3, 0))

    def file_selector(parent, title: str, description: str, var: tk.StringVar, button_text: str | None, command=None):
        box = tk.Frame(parent, bg="#FBFCFE", highlightbackground=colors["border"], highlightthickness=1)
        box.pack(fill="x", pady=(0, 12))
        left = tk.Frame(box, bg="#FBFCFE")
        left.pack(side="left", fill="x", expand=True, padx=14, pady=11)
        tk.Label(left, text=title, bg="#FBFCFE", fg=colors["navy"], font=("Segoe UI", 10, "bold"), anchor="w").pack(fill="x")
        path_label = tk.Label(left, text=shorten_path(var.get()), bg="#FBFCFE", fg=colors["muted"],
                              font=("Segoe UI", 9), anchor="w")
        path_label.pack(fill="x", pady=(3, 0))
        tk.Label(left, text=description, bg="#FBFCFE", fg="#8A98AA", font=("Segoe UI", 8), anchor="w").pack(fill="x")
        if button_text:
            ttk.Button(box, text=button_text, style="Secondary.TButton", command=command).pack(side="right", padx=12)
        var.trace_add("write", lambda *_: path_label.configure(text=shorten_path(var.get())))

    main = tk.Frame(root, bg=colors["bg"])
    main.pack(fill="both", expand=True, padx=28, pady=24)

    header = tk.Frame(main, bg=colors["navy"], height=68)
    header.pack(fill="x", pady=(0, 24))
    header.pack_propagate(False)
    header_logo = None
    try:
        header_logo = tk.PhotoImage(file=str(resource_path("app_logo.png")))
    except tk.TclError:
        pass
    if header_logo is not None:
        header_logo_widget = tk.Label(header, image=header_logo, bg=colors["navy"], borderwidth=0)
        header_logo_widget.image = header_logo
        header_logo_widget.pack(side="left", padx=(16, 12), pady=2)
    else:
        logo(header, 46).pack(side="left", padx=(16, 12), pady=11)
    brand = tk.Frame(header, bg=colors["navy"])
    brand.pack(side="left", pady=12)
    tk.Label(brand, text="RM FollowUp", bg=colors["navy"], fg="#FFFFFF", font=("Segoe UI", 16, "bold"), anchor="w").pack(anchor="w")
    tk.Label(brand, text="PO Automation  ·  Track. Follow up. Close.", bg=colors["navy"], fg="#B9D4DF", font=("Segoe UI", 9), anchor="w").pack(anchor="w")
    tk.Label(header, text="●  Runs locally", bg=colors["navy"], fg="#B8F0E9", font=("Segoe UI", 9, "bold")).pack(side="right", padx=20)

    content = tk.Frame(main, bg=colors["bg"])
    content.pack(fill="both", expand=True)
    content.grid_columnconfigure(0, weight=5, uniform="columns")
    content.grid_columnconfigure(1, weight=6, uniform="columns")
    content.grid_rowconfigure(0, weight=1)
    left = tk.Frame(content, bg=colors["bg"])
    left.grid(row=0, column=0, sticky="nsew", padx=(0, 18))
    right = tk.Frame(content, bg=colors["card"], highlightbackground=colors["border"], highlightthickness=1)
    right.grid(row=0, column=1, sticky="nsew")

    tk.Label(left, text="Daily Follow-Up Dashboard", bg=colors["bg"], fg=colors["navy"],
             font=("Segoe UI", 22, "bold"), anchor="w").pack(fill="x", pady=(8, 5))
    tk.Label(left, textvariable=dashboard_date_var, bg=colors["bg"], fg=colors["muted"],
             font=("Segoe UI", 10), anchor="w").pack(fill="x", pady=(0, 22))

    dashboard_cards = tk.Frame(left, bg=colors["bg"])
    dashboard_cards.pack(fill="x", pady=(0, 18))
    dashboard_cards.grid_columnconfigure(0, weight=1, uniform="dashboard")
    dashboard_cards.grid_columnconfigure(1, weight=1, uniform="dashboard")

    def dashboard_card(parent, column, title, value_var, accent):
        card = tk.Frame(parent, bg=colors["card"], highlightbackground=colors["border"], highlightthickness=1)
        card.grid(row=0, column=column, sticky="nsew", padx=(0 if column == 0 else 8, 8 if column == 0 else 0))
        tk.Frame(card, bg=accent, height=6).pack(fill="x")
        tk.Label(card, text=title, bg=colors["card"], fg=colors["muted"],
                 font=("Segoe UI", 10, "bold"), anchor="w").pack(fill="x", padx=18, pady=(18, 4))
        tk.Label(card, textvariable=value_var, bg=colors["card"], fg=colors["navy"],
                 font=("Segoe UI", 28, "bold"), anchor="w").pack(fill="x", padx=18, pady=(0, 18))

    dashboard_card(dashboard_cards, 0, "No Confirmation Follow-Up", no_confirmation_total_var, colors["amber"])
    dashboard_card(dashboard_cards, 1, "Delivery Follow-Up", delivery_total_var, colors["teal"])
    tk.Label(left, textvariable=dashboard_message_var, bg=colors["bg"], fg=colors["muted"],
             font=("Segoe UI", 9), justify="left", wraplength=520, anchor="w").pack(fill="x", pady=(0, 12))

    inner = tk.Frame(right, bg=colors["card"])
    inner.pack(fill="both", expand=True, padx=24, pady=24)
    tk.Label(inner, text="Workbook Setup", bg=colors["card"], fg=colors["navy"], font=("Segoe UI", 18, "bold"), anchor="w").pack(fill="x")
    tk.Label(inner, text="Choose the tracker and output location.", bg=colors["card"], fg=colors["muted"], font=("Segoe UI", 9), anchor="w").pack(fill="x", pady=(4, 18))
    file_selector(inner, "Source tracker", "Excel workbook · .xlsx or .xlsm", tracker_var, "Browse", lambda: choose_file(tracker_var, "Choose tracker workbook"))
    file_selector(inner, "Output workbook", "Generated report location", output_var, "Change", lambda: choose_file(output_var, "Choose output workbook", True))

    action_row = tk.Frame(inner, bg=colors["card"])
    action_row.pack(fill="x", pady=(4, 10))
    generate_button = ttk.Button(action_row, text="Generate Workbook", style="Primary.TButton")
    generate_button.pack(side="left", padx=(0, 8))
    drafts_button = ttk.Button(action_row, text="Generate + Create Outlook Drafts", style="Secondary.TButton")
    drafts_button.pack(side="left")
    ttk.Button(inner, text="Open Output Folder", style="Link.TButton", command=open_output_folder).pack(anchor="w", pady=(0, 14))

    status_panel = tk.Frame(inner, bg="#F7FAFC", highlightbackground=colors["border"], highlightthickness=1)
    status_panel.pack(fill="x", pady=(2, 0))
    status_dot = tk.Frame(status_panel, width=10, height=10, bg=colors["teal"])
    status_dot.pack(side="left", padx=(14, 8), pady=15)
    status_dot.pack_propagate(False)
    status_text = tk.Frame(status_panel, bg="#F7FAFC")
    status_text.pack(side="left", fill="x", expand=True, pady=10)
    tk.Label(status_text, textvariable=status_title_var, bg="#F7FAFC", fg=colors["navy"], font=("Segoe UI", 10, "bold"), anchor="w").pack(fill="x")
    tk.Label(status_text, textvariable=status_message_var, bg="#F7FAFC", fg=colors["muted"], font=("Segoe UI", 9), anchor="w", wraplength=420).pack(fill="x", pady=(2, 0))
    progress = ttk.Progressbar(inner, style="Modern.Horizontal.TProgressbar", mode="indeterminate")
    details_button = ttk.Button(inner, text="Activity details", style="Link.TButton")
    details_label = tk.Label(inner, textvariable=details_var, bg=colors["card"], fg=colors["muted"], font=("Segoe UI", 8), justify="left", anchor="w", wraplength=500)

    def toggle_details():
        if details_label.winfo_ismapped():
            details_label.pack_forget()
        else:
            details_label.pack(fill="x", pady=(4, 0))

    details_button.configure(command=toggle_details)
    details_button.pack(anchor="w", pady=(10, 0))

    def execute(with_drafts: bool):
        tracker_path = Path(tracker_var.get()).expanduser() if tracker_var.get() else None
        contacts_path = Path(contacts_var.get()).expanduser() if contacts_var.get() else None
        if not tracker_path or not contacts_path:
            set_status("Files required", "Choose both the tracker and vendor contact workbooks.", "warning")
            messagebox.showwarning("Missing files", "Choose both workbooks first.")
            return
        if not tracker_path.exists() or tracker_path.suffix.lower() not in {".xlsx", ".xlsm"}:
            set_status("Invalid tracker", "Select an existing .xlsx or .xlsm tracker workbook.", "error")
            messagebox.showerror("Invalid tracker", "The tracker file does not exist or is not an Excel workbook.")
            return
        if not contacts_path.exists() or contacts_path.suffix.lower() not in {".xlsx", ".xlsm"}:
            set_status("Invalid contact list", "Select an existing .xlsx or .xlsm contact workbook.", "error")
            messagebox.showerror("Invalid contact list", "The vendor contact file does not exist or is not an Excel workbook.")
            return
        output_path = Path(output_var.get()).expanduser()
        if output_path.suffix.lower() not in {".xlsx", ".xlsm"}:
            output_path = output_path.with_suffix(".xlsx")
            output_var.set(str(output_path))
        if output_path.exists():
            overwrite = messagebox.askyesno("Output workbook already exists", f"{output_path.name} already exists.\n\nChoose Yes to replace it. Close the workbook in Excel first.\nChoose No to create a new timestamped workbook.")
            if not overwrite:
                stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                output_path = output_path.with_name(f"{output_path.stem}_{stamp}{output_path.suffix}")
                output_var.set(str(output_path))
        try:
            set_status("Processing", "Reading workbooks and building the report…", "processing")
            progress.pack(fill="x", pady=(10, 0))
            progress.start(12)
            generate_button.state(["disabled"])
            drafts_button.state(["disabled"])
            root.update_idletasks()
            result = generate(tracker_path, contacts_path, output_path, with_drafts)
            progress.stop(); progress.pack_forget()
            set_status("Completed", f"{result['no_confirmation']} no-confirmation and {result['delivery']} delivery follow-ups prepared.", "success")
            details_var.set(f"Output: {result['output']}\nOutlook drafts created: {result['drafts_created']}  ·  Missing contacts: {result['drafts_skipped']}")
            messagebox.showinfo("Completed", f"Workbook created:\n{result['output']}\n\nOutlook drafts created: {result['drafts_created']}\nMissing contacts: {result['drafts_skipped']}")
        except PermissionError:
            progress.stop(); progress.pack_forget()
            set_status("Output unavailable", "The selected workbook is locked or unavailable.", "warning")
            alternative = output_path.with_name(f"{output_path.stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{output_path.suffix}")
            use_alternative = messagebox.askyesno("Cannot write output workbook", f"Windows denied access to the selected workbook.\n\nIt may be open in Excel, Preview, or another process.\n\nCreate a new file instead?\n{alternative.name}")
            if use_alternative:
                try:
                    result = generate(tracker_path, contacts_path, alternative, with_drafts)
                    output_var.set(str(alternative))
                    set_status("Completed", f"New workbook created: {alternative.name}", "success")
                    messagebox.showinfo("Completed", f"New workbook created:\n{result['output']}")
                except Exception as alternative_error:
                    set_status("Processing failed", str(alternative_error), "error")
                    messagebox.showerror("RM Tracker Automation", str(alternative_error))
        except Exception as exc:
            progress.stop(); progress.pack_forget()
            set_status("Processing failed", str(exc), "error")
            details_var.set(str(exc))
            messagebox.showerror("RM Tracker Automation", str(exc))
        finally:
            generate_button.state(["!disabled"])
            drafts_button.state(["!disabled"])

    generate_button.configure(command=lambda: execute(False))
    drafts_button.configure(command=lambda: execute(True))
    tracker_var.trace_add("write", refresh_dashboard)
    refresh_dashboard()
    LOGGER.info("Tkinter GUI initialized; entering mainloop")
    root.mainloop()


def run_gui() -> None:
    """Launch the redesigned frontend with the existing backend callbacks."""
    from rm_tracker_ui import run_gui as run_frontend

    run_frontend(
        generate=generate,
        read_tracker=read_tracker,
        build_followups=build_followups,
        logger=LOGGER,
        script_path=Path(__file__),
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="RM Tracker Automation")
    parser.add_argument("--tracker", type=Path)
    parser.add_argument("--contacts", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--drafts", action="store_true")
    args = parser.parse_args()
    if args.tracker and args.contacts:
        output = args.output or Path.cwd() / "RM_Tracker_Pivot_Analysis_Final.xlsx"
        print(generate(args.tracker, args.contacts, output, args.drafts))
    else:
        run_gui()


if __name__ == "__main__":
    main()
